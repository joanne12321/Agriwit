"""
Build and audit the project copywriting dataset.

Usage:
    python dataset_tools.py --input "../数据集_Sheet1 (1)(1).xlsx"

Outputs:
    dataset_corpus.json
    dataset_quality_report.json
"""
from __future__ import annotations

import argparse
import json
import re
import random
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any

from dataset_rag import normalize_platform, normalize_text, split_tags


BASE_DIR = Path(__file__).resolve().parent


FIELD_MAP = {
    "平台": "platform",
    "标题": "title",
    "正文": "body",
    "商品类别": "category",
    "商品名": "product_name",
    "价格区间": "price_tier",
    "目标受众": "audiences",
    "适用节日": "festivals",
    "宣传角度": "angles",
    "备注": "notes",
    "父记录": "parent_id",
}

REQUIRED_FIELDS = ["platform", "title", "body", "category", "product_name", "angles"]


def _load_xlsx_rows(path: Path) -> list[dict[str, Any]]:
    try:
        import pandas as pd

        df = pd.read_excel(path).fillna("")
        return df.to_dict(orient="records")
    except ImportError:
        pass
    except Exception:
        pass

    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("缺少 pandas/openpyxl，请先安装 requirements.txt 中的数据集依赖。") from exc

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [normalize_text(h) for h in rows[0]]
    result = []
    for raw in rows[1:]:
        item = {}
        for header, value in zip(headers, raw):
            if header:
                item[header] = value
        if any(normalize_text(v) for v in item.values()):
            result.append(item)
    return result


def normalize_record(raw: dict[str, Any], index: int) -> dict[str, Any]:
    record = {"id": f"copy_{index:04d}"}
    for cn, en in FIELD_MAP.items():
        value = raw.get(cn, "")
        if en in {"audiences", "festivals", "angles"}:
            record[en] = split_tags(value)
        elif en == "platform":
            record[en] = normalize_platform(value)
        else:
            record[en] = normalize_text(value)

    record["body_length"] = len(record.get("body", ""))
    record["title_length"] = len(record.get("title", ""))
    record["quality_flags"] = quality_flags(record)
    record["created_at"] = datetime.now().date().isoformat()
    return record


def quality_flags(record: dict[str, Any]) -> list[str]:
    flags: list[str] = []
    for field in REQUIRED_FIELDS:
        if not record.get(field):
            flags.append(f"missing_{field}")
    body = record.get("body", "")
    if len(body) < 50:
        flags.append("body_too_short")
    if len(body) > 800:
        flags.append("body_too_long_for_single_chunk")
    if not record.get("audiences"):
        flags.append("missing_audience")
    if not record.get("festivals"):
        flags.append("missing_festival")
    if re.search(r"(最|第一|唯一|绝对|100%)", body):
        flags.append("possible_ad_law_risk")
    return flags


def build_quality_report(records: list[dict[str, Any]]) -> dict[str, Any]:
    platform_counts = Counter(r.get("platform", "") for r in records)
    category_counts = Counter(r.get("category", "") for r in records)
    price_counts = Counter(r.get("price_tier", "") for r in records)
    festival_counts = Counter(f for r in records for f in r.get("festivals", []))
    flag_counts = Counter(flag for r in records for flag in r.get("quality_flags", []))

    body_lengths = sorted(r.get("body_length", 0) for r in records)
    median_body = body_lengths[len(body_lengths) // 2] if body_lengths else 0

    duplicate_bodies = len(records) - len({r.get("body", "") for r in records})
    ready_records = sum(1 for r in records if not r.get("quality_flags"))

    recommendations = []
    if len(records) < 300:
        recommendations.append("数据量不足：建议暑期先完成300-500条高质量样本，再扩展到1000+。")
    if platform_counts and max(platform_counts.values()) / max(1, len(records)) > 0.6:
        recommendations.append("平台分布过于集中：至少覆盖小红书、抖音/快手、电商详情页、百科/产地介绍。")
    if category_counts and max(category_counts.values()) / max(1, len(records)) > 0.45:
        recommendations.append("品类分布过于集中：单一品类建议不超过总量30%-35%。")
    if flag_counts.get("missing_audience", 0):
        recommendations.append("目标受众需要拆分为身份、地域、圈层，避免混在一个字段里。")
    if duplicate_bodies:
        recommendations.append("存在正文重复，需要去重或保留父记录关系。")
    if not recommendations:
        recommendations.append("当前数据可作为RAG语料继续扩展，并建议补充独立评价集。")

    return {
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "total_records": len(records),
        "ready_records": ready_records,
        "median_body_length": median_body,
        "duplicate_bodies": duplicate_bodies,
        "platform_counts": dict(platform_counts),
        "category_counts": dict(category_counts),
        "price_counts": dict(price_counts),
        "festival_counts": dict(festival_counts),
        "quality_flag_counts": dict(flag_counts),
        "recommendations": recommendations,
    }


def build_dataset(input_path: Path, output_dir: Path) -> tuple[Path, Path]:
    raw_rows = _load_xlsx_rows(input_path)
    records = [normalize_record(row, idx + 1) for idx, row in enumerate(raw_rows)]
    report = build_quality_report(records)

    corpus_path = output_dir / "dataset_corpus.json"
    report_path = output_dir / "dataset_quality_report.json"

    corpus = {
        "schema_version": "agriwit-copywriting-v1",
        "source_file": str(input_path),
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "records": records,
    }
    corpus_path.write_text(json.dumps(corpus, ensure_ascii=False, indent=2), encoding="utf-8")
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return corpus_path, report_path


def _sft_user_prompt(record: dict[str, Any]) -> str:
    """把结构化筛选条件整理成模型可学习的用户指令。"""
    fields = [
        ("平台", record.get("platform", "")),
        ("商品类别", record.get("category", "")),
        ("商品名", record.get("product_name", "")),
        ("价格区间", record.get("price_tier", "")),
        ("目标受众", "、".join(record.get("audiences") or [])),
        ("适用节日", "、".join(record.get("festivals") or [])),
        ("宣传角度", "、".join(record.get("angles") or [])),
        ("备注", record.get("notes", "")),
    ]
    lines = ["请根据以下信息生成农产品营销文案。"]
    lines.extend(f"{label}：{value}" for label, value in fields if value)
    lines.extend([
        "",
        "请输出：",
        "1. 标题",
        "2. 正文",
        "3. 标签",
        "要求：围绕真实产品信息创作，不要照抄示例，不要编造未提供的认证、奖项、功效或承诺。",
    ])
    return "\n".join(lines)


def _sft_assistant_answer(record: dict[str, Any]) -> str:
    title = record.get("title", "").strip()
    body = record.get("body", "").strip()
    tags = " ".join(record.get("angles") or [])
    answer = [f"标题：{title}", "", f"正文：{body}"]
    if tags:
        answer.extend(["", f"标签：{' '.join('#' + tag.lstrip('#') for tag in record['angles'][:8])}"])
    return "\n".join(answer)


def build_sft_files(
    corpus_path: Path,
    output_dir: Path,
    eval_ratio: float = 0.1,
    seed: int = 42,
) -> tuple[Path, Path, dict[str, int]]:
    """从清洗后的 corpus 生成 Qwen/LLaMA-Factory 可用的 JSONL 文件。"""
    data = json.loads(corpus_path.read_text(encoding="utf-8"))
    records = data.get("records", data) if isinstance(data, (dict, list)) else []
    eligible = [
        r for r in records
        if isinstance(r, dict)
        and r.get("title")
        and r.get("body")
        and r.get("platform")
        and r.get("product_name")
        and r.get("category")
        and not r.get("quality_flags")
    ]

    random.Random(seed).shuffle(eligible)
    eval_size = max(1, round(len(eligible) * eval_ratio)) if len(eligible) > 1 else 0
    eval_records = eligible[:eval_size]
    train_records = eligible[eval_size:]

    system_prompt = (
        "你是云稼集 Agriwit 的农产品营销文案助手。"
        "你需要根据平台、商品信息、目标受众和宣传角度，"
        "生成真实、自然、适合发布的中文营销文案。"
    )

    def to_line(record: dict[str, Any]) -> str:
        item = {
            "messages": [
                {"role": "system", "content": system_prompt},
                {"role": "user", "content": _sft_user_prompt(record)},
                {"role": "assistant", "content": _sft_assistant_answer(record)},
            ],
            "metadata": {
                "source_id": record.get("id", ""),
                "platform": record.get("platform", ""),
                "category": record.get("category", ""),
                "product_name": record.get("product_name", ""),
            },
        }
        return json.dumps(item, ensure_ascii=False)

    output_dir.mkdir(parents=True, exist_ok=True)
    train_path = output_dir / "sft_train.jsonl"
    eval_path = output_dir / "sft_eval.jsonl"
    train_path.write_text(
        "\n".join(to_line(r) for r in train_records) + ("\n" if train_records else ""),
        encoding="utf-8",
    )
    eval_path.write_text(
        "\n".join(to_line(r) for r in eval_records) + ("\n" if eval_records else ""),
        encoding="utf-8",
    )
    return train_path, eval_path, {
        "eligible": len(eligible),
        "train": len(train_records),
        "eval": len(eval_records),
        "excluded": len(records) - len(eligible),
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", default="../数据集_Sheet1 (1)(1).xlsx")
    parser.add_argument("--output-dir", default=str(BASE_DIR))
    parser.add_argument(
        "--export-sft",
        action="store_true",
        help="同时导出 Qwen/LLaMA-Factory 可用的 sft_train.jsonl 和 sft_eval.jsonl",
    )
    parser.add_argument("--eval-ratio", type=float, default=0.1)
    parser.add_argument("--seed", type=int, default=42)
    args = parser.parse_args()

    corpus_path, report_path = build_dataset(Path(args.input), Path(args.output_dir))
    print(f"dataset_corpus={corpus_path}")
    print(f"quality_report={report_path}")
    if args.export_sft:
        train_path, eval_path, stats = build_sft_files(
            corpus_path,
            Path(args.output_dir),
            eval_ratio=args.eval_ratio,
            seed=args.seed,
        )
        print(f"sft_train={train_path}")
        print(f"sft_eval={eval_path}")
        print(f"sft_stats={json.dumps(stats, ensure_ascii=False)}")


if __name__ == "__main__":
    main()
